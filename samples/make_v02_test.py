#!/usr/bin/env python3
"""후원페이지 기획서 v0.1 → v0.2 테스트 버전 만들기 (버전 비교 기능 검증용).

실제 기획 수정에서 흔한 변경을 골고루 넣는다:
  - 텍스트 수정 (8p 5초→3초, 9p 480p→540p, 20p 갱신 주기, 21p 공유 버튼)
  - 행 추가 (21p 10. 신고 버튼) + 목업에 버튼·마커 추가
  - 행 삭제 (14p 3. 영화관 모드 > 전체 화면) + 마커 삭제
  - 도형 이동 (13p PIP 창 위치)
  - 페이지 추가 (방송 종료 뒤 '네트워크 오류')
  - 페이지 삭제 (5p 영상 리스트 > 투네이션 로고 신규 추가)
  - 표지 버전, 개정 이력 추가
사용: python samples/make_v02_test.py <v0.1.pptx> <v0.2.pptx>
"""
import copy
import sys
from pptx import Presentation

NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
A = '{%s}' % NS['a']

src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src)
S = list(prs.slides)
SW = prs.slide_width


def tables(slide):
    return [sh for sh in slide.shapes if getattr(sh, 'has_table', False) and sh.has_table]


def replace_text(root, old, new, count=1):
    n = 0
    for t in root.iter(A + 't'):
        if t.text and old in t.text:
            t.text = t.text.replace(old, new)
            n += 1
            if n >= count:
                break
    assert n, f'not found: {old}'


def set_cell_lines(cell, lines):
    """셀 첫 문단을 [(text, bold), ...] 줄로 다시 쓴다 (원래 run 서식 유지)."""
    p = cell.text_frame.paragraphs[0]._p
    runs = p.findall(A + 'r')
    bold_r = next((r for r in runs if (r.find(A + 'rPr') is not None and r.find(A + 'rPr').get('b') == '1')), runs[0])
    norm_r = next((r for r in runs if not (r.find(A + 'rPr') is not None and r.find(A + 'rPr').get('b') == '1')), runs[-1])
    br_proto = p.find(A + 'br')
    for ch in list(p):
        if ch.tag in (A + 'r', A + 'br', A + 'fld'):
            p.remove(ch)
    end = p.find(A + 'endParaRPr')
    for i, (text, bold) in enumerate(lines):
        if i:
            br = copy.deepcopy(br_proto) if br_proto is not None else p.makeelement(A + 'br', {})
            (end.addprevious(br) if end is not None else p.append(br))
        r = copy.deepcopy(bold_r if bold else norm_r)
        r.find(A + 't').text = text
        (end.addprevious(r) if end is not None else p.append(r))
    for extra in cell.text_frame.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


# ---- 표지 / 개정 이력 ----
replace_text(S[0].shapes._spTree, '0.1', '0.2')
replace_text(S[0].shapes._spTree, '2026-09-30', '2026-10-01')
rev = tables(S[1])[0].table
for c_src, c_dst, text in zip(rev.rows[1].cells, rev.rows[2].cells,
                              ['0.2', '2026-10-01', '플레이어 기능 표시 시간 변경(5초→3초), 신고 버튼 추가, 네트워크 오류 화면 추가, 영상 리스트 로고 페이지 삭제', '정원영']):
    c_dst.text_frame._txBody.remove(c_dst.text_frame.paragraphs[0]._p)
    p = copy.deepcopy(c_src.text_frame.paragraphs[0]._p)
    c_dst.text_frame._txBody.append(p)
    ts = list(p.iter(A + 't'))
    ts[0].text = text
    for t in ts[1:]:
        t.text = ''

# ---- 8p: 기능 표시 시간 5초 → 3초 ----
replace_text(tables(S[7])[0]._element, ', 5', ', 3')

# ---- 9p: 480p → 540p ----
replace_text(tables(S[8])[0]._element, '② 480p', '② 540p')

# ---- 13p: PIP 창을 왼쪽 위로 이동 ----
for sh in S[12].shapes:
    if sh.name == '그룹 16':
        sh.left = sh.left - int(SW * 0.03)
        sh.top = sh.top - int(prs.slide_height * 0.04)

# ---- 14p: 3번 행(영화관 모드 > 전체 화면) + 마커 3 삭제 ----
t14 = tables(S[13])[0]
tr = t14.table._tbl.tr_lst[2]
t14.height = t14.height - tr.h
tr.getparent().remove(tr)
for sh in list(S[13].shapes):
    if sh.has_text_frame and sh.text_frame.text.strip() == '3' and sh.name.startswith('타원'):
        sh._element.getparent().remove(sh._element)

# ---- 20p: 시청자 수 갱신 주기 ----
replace_text(tables(S[19])[0]._element, '5~10', '10')

# ---- 21p: 9번 공유 버튼 문구 수정 + 10번 신고 버튼 행 추가 + 목업 버튼·마커 추가 ----
t21 = tables(S[20])[0]
tbl = t21.table
set_cell_lines(tbl.rows[8].cells[1], [('공유 버튼', True), ("- 클릭 시, 후원페이지 URL 복사 및 '링크가 복사되었습니다' 토스트 노출", False)])
tr9 = tbl._tbl.tr_lst[8]
tr10 = copy.deepcopy(tr9)
tr9.addnext(tr10)
t21.height = t21.height + tr9.h
row10 = tbl.rows[9]
for t in row10.cells[0]._tc.iter(A + 't'):
    t.text = '10' if t.text.strip() == '9' else t.text
set_cell_lines(row10.cells[1], [('신고 버튼', True), ('- 클릭 시, 신고 사유 선택 모달 호출', False)])
share = next(sh for sh in S[20].shapes if sh.name == 'Group 17')
fav = next(sh for sh in S[20].shapes if sh.name == 'Group 2')
btn = copy.deepcopy(share._element)
share._element.addnext(btn)
# 복사한 버튼: 즐겨찾기 왼쪽으로, 글자 '공유' → '신고'
off = btn.find('.//' + '{%s}grpSpPr' % NS['p']).find(A + 'xfrm').find(A + 'off')
dx = (fav.left - share.width - int(SW * 0.004)) - int(off.get('x'))
off.set('x', str(int(off.get('x')) + dx))
for t in btn.iter(A + 't'):
    if '공유' in (t.text or ''):
        t.text = t.text.replace('공유', '신고')
for sp in btn.iter('{%s}cNvPr' % NS['p']):
    sp.set('id', str(9000 + int(sp.get('id'))))
# 마커 10
m9 = next(sh for sh in S[20].shapes if sh.has_text_frame and sh.text_frame.text.strip() == '9' and sh.name.startswith('타원'))
m10 = copy.deepcopy(m9._element)
m9._element.addnext(m10)
x = m10.find('.//' + A + 'off')
x.set('x', str(fav.left - share.width // 2 - int(SW * 0.004)))
for t in m10.iter(A + 't'):
    t.text = '10'
m10.find('.//{%s}cNvPr' % NS['p']).set('id', '9999')

# ---- 페이지 추가: '송출 끊김'(18p)을 복사해 '네트워크 오류'를 방송 종료(19p) 뒤에 ----
base = S[17]
new = prs.slides.add_slide(base.slide_layout)
for ph in list(new.shapes):
    ph._element.getparent().remove(ph._element)
rid_map = {}
for rid, rel in base.part.rels.items():
    if 'notesSlide' in rel.reltype or 'slideLayout' in rel.reltype:
        continue
    rid_map[rid] = new.part.rels.get_or_add(rel.reltype, rel._target) if not rel.is_external else new.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
for el in base.shapes._spTree:
    if el.tag.endswith('}nvGrpSpPr') or el.tag.endswith('}grpSpPr'):
        continue
    e = copy.deepcopy(el)
    for node in e.iter():
        for k, v in list(node.attrib.items()):
            if k.startswith('{%s}' % NS['r']) and v in rid_map:
                node.set(k, rid_map[v])
    new.shapes._spTree.append(e)
replace_text(new.shapes._spTree, '송출 끊김', '네트워크 오류', count=5)
for sh in new.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith('방송 송출 상태가'):
        set_cell_lines(sh, [('네트워크 연결이 불안정합니다.', False), ('연결 상태를 확인한 뒤 다시 시도해 주세요.', False)])
    if sh.has_text_frame and sh.text_frame.text.startswith('방송 중단 시간'):
        set_cell_lines(sh, [('다시 시도', False)])
for sh in new.shapes:
    if getattr(sh, 'has_table', False) and sh.has_table:
        set_cell_lines(sh.table.rows[0].cells[1], [('라이브 방송 네트워크 오류 화면', True),
                                                     ('- 시청자 네트워크 연결이 끊겼을 경우 해당 화면으로 대체 노출함', False), ('', False),
                                                     ('- [다시 시도] 클릭 시 재연결 시도', False), ('', False),
                                                     ('- 재연결 성공 시, 실시간 방송 시점으로 재생', False)])
lst = prs.slides._sldIdLst
ids = list(lst)
new_id = ids[-1]
lst.remove(new_id)
lst.insert(19, new_id)          # 방송 종료(19p) 바로 뒤

# ---- 페이지 삭제: 5p ----
victim = list(lst)[4]
rid = victim.get('{%s}id' % NS['r'])
prs.part.drop_rel(rid)
lst.remove(victim)

prs.save(dst)
print('saved', dst, 'slides', len(prs.slides))
